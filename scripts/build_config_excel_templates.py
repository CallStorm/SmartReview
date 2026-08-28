#!/usr/bin/env python3
"""Generate blank and sample config Excel templates for admin manual entry."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "frontend" / "public" / "config-templates"
SOURCE_XLSX = Path(
    r"C:\Users\Administrator\Desktop\一建方案审核\落地脚手架\专项方案审核模版_落地脚手架final.xlsx"
)

META_FIELDS = ["方案大类", "方案名称", "适用说明", "编制人", "更新日期"]
STRUCTURE_HEADERS = ["章节编码", "章节标题", "是否校验", "完整性依据", "备注"]
CONSISTENCY_HEADERS = ["规则ID", "规则名称", "是否启用", "源章节编码", "目标章节编码", "检查说明"]
BASIS_HEADERS = [
    "依据ID",
    "文献类型",
    "标准号或文号",
    "文献名称",
    "版本或施行日期说明",
    "是否必引",
    "类别",
    "分类",
    "备注",
]
CHAPTER_HEADERS = [
    "章节编码",
    "依赖章节编码",
    "知识库引用",
    "人工提示词",
    "图审核启用",
    "有无图",
    "图种类",
    "图内容",
]
KB_HEADERS = ["条目ID", "知识库名称", "TAG信息", "内容引用", "摘要说明", "是否启用", "备注"]
SHEET_ORDER = [
    "说明与版本",
    "章节结构_完整性",
    "一致性规则",
    "编制依据库",
    "章节审查配置",
    "知识库",
]

BLANK_FILENAME = "专项方案审核配置模版_空白.xlsx"
SAMPLE_FILENAME = "专项方案审核配置模版_落地脚手架示例.xlsx"

SCAFFOLDING_TEXT_REPLACEMENTS = [
    ("基坑土方开挖、支护技术参数", "脚手架搭设参数、荷载取值"),
    ("知识库:基坑工程 tag:土钉墙支护", "知识库:脚手架工程 tag:落地脚手架"),
    ("基坑工程", "脚手架工程"),
    ("土钉墙支护", "落地脚手架"),
    ("DEEP_EXCAVATION", ""),
]


def _write_meta_sheet(ws: Worksheet, values: dict[str, str]) -> None:
    ws.delete_rows(1, ws.max_row or 1)
    ws.append(["字段名", "值"])
    for field in META_FIELDS:
        ws.append([field, values.get(field, "")])


def _delete_column_by_header(ws: Worksheet, header_name: str) -> None:
    headers = [cell.value for cell in ws[1]]
    if header_name not in headers:
        return
    ws.delete_cols(headers.index(header_name) + 1)


def _fix_scaffolding_text(value):
    if not isinstance(value, str):
        return value
    result = value
    for old, new in SCAFFOLDING_TEXT_REPLACEMENTS:
        result = result.replace(old, new)
    return result


def _fix_sheet_text(ws: Worksheet) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = _fix_scaffolding_text(cell.value)


def _reorder_sheets(wb: Workbook) -> None:
    for target_idx, name in enumerate(SHEET_ORDER):
        current_idx = wb.sheetnames.index(name)
        if current_idx != target_idx:
            wb.move_sheet(wb[name], offset=target_idx - current_idx)


def _build_blank_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    header_map = {
        "章节结构_完整性": STRUCTURE_HEADERS,
        "一致性规则": CONSISTENCY_HEADERS,
        "编制依据库": BASIS_HEADERS,
        "章节审查配置": CHAPTER_HEADERS,
        "知识库": KB_HEADERS,
    }

    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(sheet_name)
        if sheet_name == "说明与版本":
            _write_meta_sheet(
                ws,
                {
                    "方案大类": "",
                    "方案名称": "",
                    "适用说明": "请按实际方案类型填写",
                    "编制人": "",
                    "更新日期": "",
                },
            )
        else:
            ws.append(header_map[sheet_name])

    return wb


def _transform_chapter_review_sheet(ws: Worksheet) -> None:
    _delete_column_by_header(ws, "审查侧重点")
    _delete_column_by_header(ws, "数值审核")

    headers = [cell.value for cell in ws[1]]
    for col_name in ["图审核启用", "有无图", "图种类", "图内容"]:
        if col_name not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=col_name)
            headers.append(col_name)

    image_examples = {
        (1.2, "1.2"): (
            "是",
            "有",
            "施工总平面布置图",
            "核对平面布置图是否完整标注脚手架、临建、通道等要素",
        ),
        (8.3, "8.3"): (
            "是",
            "有",
            "路线图",
            "检查是否有救援路线图",
        ),
    }
    for row_idx in range(2, ws.max_row + 1):
        chapter_code = ws.cell(row=row_idx, column=1).value
        for keys, values in image_examples.items():
            if chapter_code in keys:
                for col_offset, value in enumerate(values, start=5):
                    ws.cell(row=row_idx, column=col_offset, value=value)
                break


def _build_sample_workbook() -> Workbook:
    wb = load_workbook(SOURCE_XLSX)

    if "审核流程图" in wb.sheetnames:
        del wb["审核流程图"]

    _write_meta_sheet(
        wb["说明与版本"],
        {
            "方案大类": "脚手架工程",
            "方案名称": "落地脚手架",
            "适用说明": "落地脚手架专项方案审核配置示例",
            "编制人": "",
            "更新日期": "",
        },
    )

    consistency_ws = wb["一致性规则"]
    _delete_column_by_header(consistency_ws, "严重等级")
    _fix_sheet_text(consistency_ws)

    _transform_chapter_review_sheet(wb["章节审查配置"])
    _fix_sheet_text(wb["章节审查配置"])

    for sheet_name in SHEET_ORDER:
        if sheet_name in {"说明与版本", "一致性规则", "章节审查配置"}:
            continue
        _fix_sheet_text(wb[sheet_name])

    _reorder_sheets(wb)
    return wb


def _ensure_output_dir() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def build_templates() -> tuple[Path, Path]:
    _ensure_output_dir()
    blank_path = OUTPUT_DIR / BLANK_FILENAME
    sample_path = OUTPUT_DIR / SAMPLE_FILENAME

    _build_blank_workbook().save(blank_path)
    _build_sample_workbook().save(sample_path)
    return blank_path, sample_path


if __name__ == "__main__":
    blank, sample = build_templates()
    print(f"Wrote {blank}")
    print(f"Wrote {sample}")
