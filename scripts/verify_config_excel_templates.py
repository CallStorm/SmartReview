#!/usr/bin/env python3
"""Verify generated config Excel templates meet Task 1 requirements."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "frontend" / "public" / "config-templates"

META_FIELDS = ["方案大类", "方案名称", "适用说明", "编制人", "更新日期"]
CONSISTENCY_HEADERS = ["规则ID", "规则名称", "是否启用", "源章节编码", "目标章节编码", "检查说明"]
CHAPTER_HEADERS = [
    "章节编码",
    "依赖章节编码",
    "知识库引用",
    "人工提示词",
    "图审核启用",
    "有无图",
    "图种类",
    "图内容",
]
SHEET_ORDER = [
    "说明与版本",
    "章节结构_完整性",
    "一致性规则",
    "编制依据库",
    "章节审查配置",
    "知识库",
]

BLANK_FILENAME = "专项方案审核配置模版_空白.xlsx"
SAMPLE_FILENAME = "专项方案审核配置模版_落地脚手架示例.xlsx"

FORBIDDEN_META_FIELDS = {"方案类型代码", "模版版本"}


def _fail(message: str) -> None:
    print(message, file=sys.stderr)
    sys.exit(1)


def _header_row(ws) -> list:
    return [cell.value for cell in ws[1]]


def _meta_field_map(ws) -> dict[str, str | None]:
    fields: dict[str, str | None] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        value = row[1] if len(row) > 1 else None
        fields[key] = None if value is None else str(value)
    return fields


def _scan_for_token(wb, token: str) -> list[str]:
    hits: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None and token in str(cell):
                    hits.append(f"{sheet_name}: {cell}")
    return hits


def _verify_common(wb, label: str) -> None:
    if wb.sheetnames != SHEET_ORDER:
        _fail(f"{label}: sheet order mismatch: {wb.sheetnames!r} != {SHEET_ORDER!r}")

    meta_ws = wb["说明与版本"]
    meta_fields = _meta_field_map(meta_ws)
    meta_keys = set(meta_fields)
    if meta_keys != set(META_FIELDS):
        _fail(
            f"{label}: 说明与版本 fields mismatch: {sorted(meta_keys)!r} != {META_FIELDS!r}"
        )
    forbidden = meta_keys & FORBIDDEN_META_FIELDS
    if forbidden:
        _fail(f"{label}: forbidden meta fields present: {sorted(forbidden)!r}")

    consistency_headers = _header_row(wb["一致性规则"])
    if consistency_headers != CONSISTENCY_HEADERS:
        _fail(
            f"{label}: 一致性规则 headers mismatch: {consistency_headers!r} != {CONSISTENCY_HEADERS!r}"
        )

    chapter_headers = _header_row(wb["章节审查配置"])
    if chapter_headers != CHAPTER_HEADERS:
        _fail(
            f"{label}: 章节审查配置 headers mismatch: {chapter_headers!r} != {CHAPTER_HEADERS!r}"
        )


def _verify_blank(path: Path) -> None:
    if not path.is_file():
        _fail(f"Missing blank template: {path}")
    wb = load_workbook(path, data_only=True)
    _verify_common(wb, "blank")


def _verify_sample(path: Path) -> None:
    if not path.is_file():
        _fail(f"Missing sample template: {path}")
    wb = load_workbook(path, data_only=True)
    _verify_common(wb, "sample")

    if "审核流程图" in wb.sheetnames:
        _fail("sample: 审核流程图 sheet should be removed")

    meta = _meta_field_map(wb["说明与版本"])
    if meta.get("方案大类") != "脚手架工程":
        _fail(f"sample: 方案大类 expected 脚手架工程, got {meta.get('方案大类')!r}")
    if meta.get("方案名称") != "落地脚手架":
        _fail(f"sample: 方案名称 expected 落地脚手架, got {meta.get('方案名称')!r}")

    deep_hits = _scan_for_token(wb, "DEEP_EXCAVATION")
    if deep_hits:
        _fail(f"sample: found DEEP_EXCAVATION: {deep_hits[:3]}")


def main() -> None:
    blank_path = TEMPLATE_DIR / BLANK_FILENAME
    sample_path = TEMPLATE_DIR / SAMPLE_FILENAME
    _verify_blank(blank_path)
    _verify_sample(sample_path)
    print("All config Excel template checks passed.")


if __name__ == "__main__":
    main()
