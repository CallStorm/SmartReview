"""Generate audit report Word documents from ReviewReportV1 JSON."""

from __future__ import annotations

import io
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table, _Cell

from app.models.scheme_review_task import SchemeReviewTask
from app.schemas.review_report import ReportStep, ReviewReportV1
from app.services import review_report_content as content


_CN_FONT = "宋体"
_CN_FONT_SIZES = {
    "title": Pt(16),
    "h2": Pt(13),
    "h3": Pt(11),
    "body": Pt(10.5),
    "quote": Pt(9.5),
    "cell": Pt(9),
}
_HEADER_FILL = "F0F0F0"
_BORDER_COLOR = "999999"
_BORDER_SIZE = "4"

_TABLE_STYLE = (
    "Table Grid"
)


def _set_cn_font(run, size: Pt) -> None:
    run.font.name = _CN_FONT
    rpr = run._element.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        from docx.oxml import OxmlElement

        r_fonts = OxmlElement("w:rFonts")
        rpr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), _CN_FONT)
    r_fonts.set(qn("w:ascii"), _CN_FONT)
    r_fonts.set(qn("w:hAnsi"), _CN_FONT)
    run.font.size = size


def _set_cell_text(
    cell: _Cell,
    text: str,
    *,
    bold: bool = False,
    size: Pt | None = None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for line_idx, line in enumerate(text.split("\n")):
        if line_idx == 0:
            run = p.add_run(line)
        else:
            sub_p = cell.add_paragraph()
            sub_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = sub_p.add_run(line)
        run.bold = bold
        _set_cn_font(run, size or _CN_FONT_SIZES["cell"])
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def _apply_table_style(table: Table) -> None:
    try:
        table.style = _TABLE_STYLE
    except KeyError:
        table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def _set_header_fill(table: Table) -> None:
    from docx.oxml import OxmlElement

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), _HEADER_FILL)
        tc_pr.append(shd)


def _add_paragraph(
    doc: Document,
    text: str,
    *,
    kind: str,
    align: int | None = None,
    bold: bool = False,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    _set_cn_font(run, _CN_FONT_SIZES[kind])


def _add_kv_table(
    doc: Document,
    rows: Iterable[tuple[str, str]],
) -> Table:
    rows_list = list(rows)
    table = doc.add_table(rows=1 + len(rows_list), cols=2)
    _apply_table_style(table)
    _set_header_fill(table)
    _set_cell_text(table.rows[0].cells[0], "项目", bold=True)
    _set_cell_text(table.rows[0].cells[1], "内容", bold=True)
    for idx, (label, value) in enumerate(rows_list, start=1):
        _set_cell_text(table.rows[idx].cells[0], label, bold=True)
        _set_cell_text(table.rows[idx].cells[1], value)
    return table


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> Table:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _apply_table_style(table)
    _set_header_fill(table)
    for col_idx, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[col_idx], header, bold=True)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            _set_cell_text(table.rows[row_idx].cells[col_idx], str(value))
    return table


def _render_structure(doc: Document, step: ReportStep) -> None:
    rows = content.structure_table_rows(step)
    if rows:
        _add_table(doc, ["章节", "问题类型", "说明"], rows)


def _render_compilation_basis(doc: Document, step: ReportStep) -> None:
    rows = content.compilation_basis_table_rows(step)
    if rows:
        _add_table(
            doc,
            ["问题", "类别", "章节", "原文", "修改建议", "标准依据"],
            rows,
        )


def _render_context_consistency(doc: Document, step: ReportStep) -> None:
    rows = content.context_consistency_table_rows(step)
    if rows:
        _add_table(doc, ["章节", "对比章节", "问题", "优化建议"], rows)


def _render_content_like(doc: Document, step: ReportStep) -> None:
    for chapter, rows in content.content_like_groups(step):
        _add_paragraph(doc, chapter, kind="h3")
        _add_table(doc, ["序号", "问题", "原文内容", "修改建议"], rows)


def _render_step(doc: Document, section_num: int, step: ReportStep) -> None:
    _add_paragraph(doc, content.section_heading(section_num, step), kind="h2")
    if step.summary:
        _add_paragraph(doc, f"步骤摘要:{step.summary}", kind="body")

    desc = content.STEP_DESCRIPTION.get(step.step_id, "")
    if desc:
        _add_paragraph(doc, desc, kind="quote")

    if not step.issues:
        _add_paragraph(doc, "本步骤无问题项。", kind="body")
        return

    if step.step_id == "structure":
        _render_structure(doc, step)
    elif step.step_id == "compilation_basis":
        _render_compilation_basis(doc, step)
    elif step.step_id == "context_consistency":
        _render_context_consistency(doc, step)
    else:
        _render_content_like(doc, step)


def build_audit_report_docx(
    task: SchemeReviewTask,
    report: ReviewReportV1,
    *,
    system_name: str = "智能方案审核",
) -> bytes:
    scheme_name = content.scheme_display_name(task)
    doc = Document()
    doc.core_properties.title = f"[{scheme_name}]审核报告"
    doc.core_properties.author = system_name

    _add_paragraph(
        doc,
        f"[{scheme_name}]审核报告",
        kind="title",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    _add_kv_table(doc, content.meta_kv_rows(task, report, system_name=system_name))

    for section_num, step in enumerate(content.iter_ordered_steps(report), start=1):
        _render_step(doc, section_num, step)

    _add_paragraph(doc, content.DISCLAIMER, kind="body")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()