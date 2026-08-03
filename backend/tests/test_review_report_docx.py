"""Tests for audit report Word (docx) generation."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime

from docx import Document

from app.api.review_tasks import _audit_report_docx_filename
from app.schemas.review_report import ReportIssue, ReportStep, ReviewReportV1
from app.services.review_report_docx import build_audit_report_docx


class _FakeScheme:
    category = "模板工程及支撑体系"
    name = "高大模板支撑"


class _FakeTask:
    original_filename = "睦邻中心高大支模安全专项方案测试.docx"
    status = "succeeded"
    finished_at = datetime.now(UTC)
    updated_at = datetime.now(UTC)
    scheme_type = _FakeScheme()


_CHAPTER = "一、工程概况 > 1.模板支撑体系工程概况和特点"


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    chunks: list[str] = []

    def push(text: str) -> None:
        if text:
            chunks.append(text)

    for paragraph in doc.paragraphs:
        push(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                push(cell.text)
    return "".join(chunks)


def test_build_audit_report_docx_contains_sections_and_tables() -> None:
    report = ReviewReportV1(
        steps=[
            ReportStep(
                step_id="structure",
                passed=False,
                summary="共 1 项结构问题",
                issues=[
                    ReportIssue(
                        message="缺少应急预案章节",
                        severity="error",
                        anchor={"title_path": ["八、应急处置措施"]},
                        related={"kind": "missing_section"},
                    )
                ],
            ),
            ReportStep(
                step_id="compilation_basis",
                passed=False,
                summary="发现 1 条编制依据问题",
                issues=[
                    ReportIssue(
                        message="未引用现行规范",
                        severity="error",
                        evidence="无",
                        related={
                            "category": "现行缺失",
                            "standard_no": "GB51210-2016",
                            "doc_name": "建筑施工脚手架安全技术统一标准",
                            "suggestions": ["在编制依据中补充引用"],
                        },
                    )
                ],
            ),
            ReportStep(
                step_id="content",
                passed=False,
                summary="发现 2 条内容问题",
                issues=[
                    ReportIssue(
                        message="未明确高大支模范围内梁的跨度",
                        severity="error",
                        evidence="第3节高大支模部位表格中仅列出梁截面尺寸",
                        anchor={"title_path": _CHAPTER.split(" > ")},
                        related={
                            "suggestions": [
                                "在表格或正文中补充梁的跨度数据",
                            ],
                        },
                    ),
                    ReportIssue(
                        message="未描述周边环境情况",
                        severity="error",
                        evidence="第4节仅说明高大模板立杆支承在基础和楼板上",
                        anchor={"title_path": _CHAPTER.split(" > ")},
                        related={
                            "suggestions": [
                                "补充支撑地基的承载力、压实度等参数",
                            ],
                        },
                    ),
                ],
            ),
            ReportStep(
                step_id="context_consistency",
                passed=True,
                summary="无问题",
                issues=[],
            ),
        ]
    )
    data = build_audit_report_docx(_FakeTask(), report, system_name="智能方案审核")

    assert data.startswith(b"PK")
    assert len(data) > 2000

    doc = Document(io.BytesIO(data))
    assert doc.core_properties.title is not None
    assert "审核报告" in (doc.core_properties.title or "")

    texts = _docx_text(data)
    assert "审核报告" in texts
    assert "结构审核" in texts
    assert "编制依据审核" in texts
    assert "内容审核" in texts
    assert "缺少应急预案章节" in texts
    assert "未引用现行规范" in texts
    assert "未明确高大支模范围内梁的跨度" in texts
    assert "本步骤无问题项" in texts
    assert "方案类型" in texts
    assert "摘要" in texts
    assert "本次审核任务状态为" in texts
    assert "模板支撑体系工程概况和特点" in texts


def test_parse_review_report_json_and_docx_filename() -> None:
    assert _audit_report_docx_filename("test方案.docx") == "test方案_审核报告.docx"
    raw = json.dumps(
        {"version": 1, "steps": [{"step_id": "structure", "passed": True, "issues": []}]}
    )
    parsed = ReviewReportV1.model_validate_json(raw)
    assert parsed.steps[0].step_id == "structure"